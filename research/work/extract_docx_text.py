from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def extract_document(path: Path) -> str:
    document = Document(path)
    blocks: list[str] = []

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            blocks.append(f"[P{index:03d}] {text}")

    for table_index, table in enumerate(document.tables, start=1):
        blocks.append(f"\n[TABLE {table_index}]")
        for row_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
            blocks.append(f"[R{row_index:03d}] " + " | ".join(cells))

    return "\n".join(blocks) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(extract_document(args.input), encoding="utf-8")


if __name__ == "__main__":
    main()
