from __future__ import annotations

from datetime import date
from html import escape
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
HTML_PATH = OUT_DIR / "geocheck-business-plan.html"
DOCX_PATH = OUT_DIR / "GeoCheck商業企劃書.docx"

LABELS = {
    "F": ("事實", "fact"),
    "O": ("觀察", "observed"),
    "I": ("推論", "inference"),
    "H": ("假設", "hypothesis"),
    "U": ("未知", "unknown"),
    "R": ("建議", "recommendation"),
}

SOURCES = [
    ("S01", "GeoCheck Project Charter", "本地專案文件", "docs/PROJECT_CHARTER.md", "研究優先序、產品邊界、目標使用者與商業模式仍待確認。"),
    ("S02", "GeoCheck Current State", "本地專案文件", "docs/CURRENT_STATE.md", "四題單站觀測、D1 報告保存、GA4 事件與現行運作狀態。"),
    ("S03", "AI Trust Index v1.0.0", "本地演算法規格", "docs/AI_TRUST_INDEX_V1.md", "65% 回答採用與 35% 第一方來源證據；unknown 不等於零。"),
    ("S04", "GeoCheck 競品證偽與產品路線建議", "本地研究報告", "docs/research/competitive-falsification-2026-09-07/REPORT.md", "競品、市場、差異化、實驗與六個月方向。"),
    ("S05", "Agentic Workspace Idea Check", "本地驗證報告", "docs/research/agentic-workspace-idea-check-2026-09-07.md", "Dashboard 加受限 Agent 的技術與商業可行性檢查。"),
    ("S06", "Ahrefs AI adjusted volume", "官方說明", "https://help.ahrefs.com/en/articles/16755865-what-is-ai-adjusted-volume-and-how-is-it-calculated", "Google demand 乘平台 usage ratio 的需求代理方法。"),
    ("S07", "Ahrefs AI visibility metrics", "官方說明", "https://help.ahrefs.com/en/articles/15501968-ai-visibility-metrics", "AI impressions 與 share of voice 的估算方式。"),
    ("S08", "Peec AI funding and ARR announcement", "供應商公告", "https://peec.ai/blog/we-raised-21m-series-a-to-help-brands-win-in-ai-search", "自報超過 400 萬美元 ARR 與 1,300 個品牌或代理商；未經審計。"),
    ("S09", "Otterly AI GSC workflow", "官方產品文章", "https://otterly.ai/blog/otterlyai-public-api-claude-skill-launch/", "GSC 匯出、意圖分類、prompt 改寫與查詢追蹤。"),
    ("S10", "Google Generative AI performance reports", "Google 官方公告", "https://developers.google.com/search/blog/2026/06/gen-ai-performance-reports", "Search Console 提供生成式 AI impressions、pages、countries、devices 與時間維度。"),
    ("S11", "Hogiah local business solution", "供應商產品頁", "https://www.hogiah.com/zh-TW/solutions/local", "台灣餐飲與在地商家、六個 AI 平台與情境查詢定位。"),
    ("S12", "Hogiah plans", "供應商方案頁", "https://www.hogiah.com/zh-TW/plans", "台灣定價、監測週期與公開方案內容。"),
    ("S13", "Profound Prompt Volumes", "官方產品頁", "https://www.tryprofound.com/features/prompt-volumes", "授權 panel 與需求建模。"),
    ("S14", "AEO Mantis pricing", "供應商方案頁", "https://www.aeomantis.com/zh-Hant/pricing", "低價自助監測方案與配額。"),
    ("S15", "Semrush AI Visibility Toolkit", "官方說明", "https://www.semrush.com/kb/1493-ai-visibility-toolkit", "SEO Suite 內建 AI visibility、topic demand 與 audit。"),
    ("S16", "Thunderfront restaurant plans", "供應商方案頁", "https://www.thunderfront.co/", "公開餐飲網站與在地 SEO 月費方案；只能證明供應存在。"),
    ("S17", "Jackopa restaurant marketing", "供應商方案頁", "https://jackopa.com/industry/restaurant/", "餐飲 Google 商家、內容與月費維護方案；不能證明銷量。"),
    ("S18", "Link Node restaurant pricing", "供應商方案頁", "https://link-node.com/pricing", "餐飲官網、訂位、點餐與曝光模組的低價替代。"),
    ("S19", "Taiwan accommodation and food service census", "主計總處普查", "https://www.stat.gov.tw/public/Attachment/492995126GFZJEAGL.pdf", "歷史普查顯示餐飲企業數量龐大；不是 GeoCheck 可服務市場規模。"),
    ("S20", "OneGlanse", "公開原始碼", "https://github.com/aryamantodkar/oneglanse", "可自架 AI visibility tracker 的替代方案。"),
]


SECTIONS = [
    {
        "id": "decision",
        "eyebrow": "01 Investment decision",
        "title": "投資判斷與決策條件",
        "blocks": [
            ("callout", "R", "建議為 Modify", "保留 GeoCheck 的可追溯觀測能力，把產品方向改為餐飲垂直的顧問工作台。先投資六個月驗證，不先投資完整 SaaS 平台。六個月後只依付費、第二次使用、交付工時與毛利決定是否擴大。"),
            ("p", "F", "AI visibility 已有付費市場。Peec 的公司公告自報超過 400 萬美元 ARR，Ahrefs、Semrush、Similarweb 與多家獨立 SaaS 也已提供相關產品。[S08][S15] 這支持市場存在，但不證明台灣小型顧問會付費。"),
            ("p", "I", "GeoCheck 的機會不在於再做一個顯示 mention rate 的 Dashboard，而在於替顧問完成每月最耗判斷力的工作：指出哪個可見度變化值得處理、證據在哪裡、下一步要檢查什麼，以及何時回看結果。"),
            ("table", ["投資問題", "目前判斷", "證據強度", "決策影響"], [
                ["市場是否存在", "存在國際付費類別", "中", "可繼續驗證，不應直接 Stop"],
                ["台灣 TA 是否付費", "尚不能判斷", "低", "不得先建完整帳號、計費與整合"],
                ["Agent 是否形成差異", "產品形態較好，但不是護城河", "中", "Agent 必須受限於證據與固定任務"],
                ["餐飲是否適合起步", "適合做垂直素材，不代表最大市場", "中低", "餐飲是第一套 playbook，不是永久市場邊界"],
                ["是否適合一至兩人", "適合小規模產品化服務與窄 SaaS", "中", "以人工輔助、低 API 成本與清楚停損為前提"],
            ]),
            ("kpis", [
                ("6 個月", "建議投資期限", "先買證據，不買規模"),
                ("2 次", "必要重複使用", "同一 Project 要有第二輪"),
                ("≥70%", "目標直接毛利", "通過產品化後再追求"),
                ("0", "未知資料補零", "研究與產品共同禁區"),
            ]),
        ],
    },
    {
        "id": "starting-point",
        "eyebrow": "02 Current position",
        "title": "產品起點與不可越過的邊界",
        "blocks": [
            ("p", "F", "GeoCheck 現在是一個單站、四個非品牌 discovery queries 的可觀測檢測。AI Trust Index v1 由回答採用率 65% 與已驗證第一方來源證據率 35% 組成；無有效答案時輸出 unknown，而不是零。[S02][S03]"),
            ("p", "F", "目前的 Project、長期 Monitoring、GSC onboarding、Agent Workspace 與訂閱商業模式都還不是已驗證能力。專案 Charter 仍把研究白皮書放在能力展示與商業化之前，且商業模式沒有正式定案。[S01]"),
            ("flow", [
                ("現在", "URL 檢測", "四題觀測與來源證據"),
                ("下一步", "Evidence Workspace", "保存 Project、差異與決策"),
                ("有條件", "Decision Agent", "限定問題、引用證據、人工批准"),
                ("最後", "Subscription", "只有重複使用與付費成立才開啟"),
            ]),
            ("callout", "R", "產品承諾", "對外承諾應維持在可觀測證據、診斷與決策協助。未完成因果驗證前，不承諾 AI 可見度會帶來營收，也不把需求代理值包裝成真實 AI 曝光人口。"),
        ],
    },
    {
        "id": "market",
        "eyebrow": "03 Market reality",
        "title": "市場現實與類別成熟度",
        "blocks": [
            ("p", "F", "國際市場已跨過純概念期。企業產品、獨立 SaaS、SEO Suite 與開源 tracker 同時存在；基礎監測、prompt 管理、share of voice 與報表正在商品化。[S04][S14][S15][S20]"),
            ("p", "F", "Google 已在 Search Console 提供生成式 AI 搜尋的 impressions、pages、countries、devices 與時間維度。[S10] 這會降低單純提供 Google AI 曝光圖表的價值，但仍未解決跨引擎比較、顧問交付、來源證據與下一步決策。"),
            ("p", "I", "市場階段可描述為商業化已成立、功能快速擴張、基礎監測商品化。現有資料不足以計算台灣市場規模、市占或續訂率，也不足以宣稱產業已進入整併。"),
            ("table", ["力量", "對市場的影響", "GeoCheck 應對"], [
                ["SEO Suite 內建", "買家可在原工具取得基本 AI visibility", "不要靠圖表與 prompt 數量收費"],
                ["Google 第一方資料", "免費基礎曝光資料增加", "用於校驗與上下文，不與 Google 比資料所有權"],
                ["低價 SaaS 與開源", "監測成本與價格透明", "限制監測範圍，以決策交付取代無限掃描"],
                ["台灣在地競品", "繁中、LINE、餐飲案例與在地通路已存在", "不複製內容量與引擎數競賽"],
            ]),
        ],
    },
    {
        "id": "ta",
        "eyebrow": "04 Target customer",
        "title": "建議目標客群與排除條件",
        "blocks": [
            ("callout", "H", "建議 TA", "主要付費者是有 3 至 15 個長期客戶、每月需要交付 SEO、內容或曝光決策的獨立顧問與小型行銷團隊。餐飲品牌與分店是被管理的 Project，不是初期主要買家。"),
            ("p", "O", "台灣公開市場同時存在一次性網站案與月費型餐飲曝光服務。月費方案通常包含 Google 商家、內容、評論、SEO 或廣告執行，而不是單獨購買監測。[S16][S17][S18] 這支持持續工作存在，也顯示純監測的支付意願可能很弱。"),
            ("table", ["區隔", "痛點頻率", "付費理由", "判斷"], [
                ["單店餐廳老闆", "低至事件型", "希望有客人上門", "不作核心 TA；可作免費健檢入口"],
                ["一次性網站工作室", "專案結束後低", "準時交付網站", "排除"],
                ["有月約的 SEO 或內容顧問", "每月", "節省分析與月報時間，提高交付可信度", "第一優先"],
                ["管理多據點餐飲品牌的內部團隊", "每週至每月", "品牌與分店資訊一致、跨區域追蹤", "第二階段"],
                ["大型 Agency 或 Enterprise", "持續", "權限、API、SSO、採購與服務", "初期排除，銷售成本過高"],
            ]),
            ("p", "U", "目前沒有台灣合格顧問母體數、可接觸名單、採購訪談或實際付費資料，因此不能宣稱這個 TA 足以支撐大型公司。合理目標是先驗證能否支撐創辦人型 micro SaaS。"),
        ],
    },
    {
        "id": "problem",
        "eyebrow": "05 Job to be done",
        "title": "顧客工作與付費價值",
        "blocks": [
            ("p", "H", "顧問真正需要完成的工作不是知道一個分數，而是每月向客戶回答四件事：發生了什麼、為什麼值得處理、這期先做哪一件事、下次如何證明有沒有改善。"),
            ("table", ["目前做法", "成本或缺口", "GeoCheck 提案"], [
                ["手動測多個 AI 引擎", "結果易變、難追溯、成本不穩", "固定少量 sentinel queries 與完整 run 證據"],
                ["截圖與試算表", "版本分散，月報需重組", "Project history 與可引用的 evidence IDs"],
                ["通用 AI 摘要報表", "容易生成無證據建議", "每個判斷附來源、信心與限制"],
                ["只看 GSC 或 GA4", "涵蓋單一平台或只有到站後行為", "分開呈現，不混成虛構 ROI 分數"],
            ]),
            ("callout", "R", "北極星價值", "每個月替顧問省下一次資料整理與判讀，並產生至少一項能被客戶採用、且能在下一輪回查的決策。若只得到漂亮摘要，產品沒有成立。"),
        ],
    },
    {
        "id": "solution",
        "eyebrow": "06 Product strategy",
        "title": "半 Agent 加 Dashboard 的產品形態",
        "blocks": [
            ("p", "R", "採用一個 Evidence Workspace，而不是多 Agent 平台。Dashboard 負責保存觀測事實；Agent 只回答有限決策問題；GeoBooks 保存經人確認後可重跑的流程；Decision Ledger 記錄採用、延後、拒絕與下次驗證。"),
            ("flow", [
                ("Project", "品牌與分店", "網站、題組、版本、資料權限"),
                ("Evidence", "可觀測證據", "答案、引用、官方 URL、unknown"),
                ("Diagnose", "受限 Agent", "解釋變化、找來源缺口、標信心"),
                ("Decide", "人工批准", "接受、拒絕、延後、指定負責人"),
                ("Verify", "下一輪回查", "同一決策與新證據比較"),
            ]),
            ("table", ["模組", "第一版能力", "明確不做"], [
                ["Dashboard", "Project、原始觀測、差異、來源與有效分母", "不做全網真實市占宣稱"],
                ["Geo Agent", "解釋一個變化、提出有限檢查清單", "不自由瀏覽客戶帳號、不自動改站"],
                ["GeoBooks", "人完成一次流程後才保存與重跑", "不讓模型自行發明流程"],
                ["Decision Ledger", "記錄理由、證據、信心、狀態與回查日期", "不把建議當成已執行或已產生效果"],
            ]),
            ("p", "I", "Ploy 類型的 workspace 與 playbook 互動值得借鏡，但產品寬度不適合一至兩人複製。Ahrefs 等產品也已有讀取 workspace 並回答問題的 Agent，因此聊天介面本身不構成護城河。[S05]"),
        ],
    },
    {
        "id": "moat",
        "eyebrow": "07 Defensibility",
        "title": "差異化與可累積資產",
        "blocks": [
            ("p", "I", "可能形成防禦力的不是模型或 Agent，而是餐飲實體資料品質、可追溯決策歷史、顧問反覆採用的工作流程，以及對 unknown 與品牌分店混淆的保守處理。這些資產只有在真實重複使用中才會累積。"),
            ("table", ["候選資產", "如何累積", "失效條件"], [
                ["餐飲 entity schema", "品牌、分店、地址、菜單、價位、訂位與官方來源版本", "只有公開欄位，競品可快速複製"],
                ["Evidence quality rules", "品牌分店消歧、第一方來源驗證、unknown 規則", "判斷不比通用模型穩定"],
                ["Decision history", "建議、採用、執行證據與下一輪結果", "使用者不回來或不記錄"],
                ["Restaurant GeoBooks", "把成功的月報與檢查流程重跑", "流程只是一般 SEO checklist"],
                ["Distribution corpus", "公開案例、研究與方法說明持續帶入合格流量", "內容沒有合格讀者或只吸引免費使用者"],
            ]),
            ("callout", "U", "護城河尚未成立", "目前沒有長期客戶資料、回饋閉環、獨家需求資料或留存歷史。企劃書描述的是資產形成機制，不是既有優勢。"),
        ],
    },
    {
        "id": "business-model",
        "eyebrow": "08 Business model",
        "title": "商業模式與定價實驗",
        "blocks": [
            ("p", "R", "先以產品化服務驗證，再轉窄 SaaS。創辦人親自完成前三個 Project 的月度決策包，以此找出真正重複的步驟；只有重複部分進入產品。"),
            ("table", ["層級", "範圍", "價格假設", "驗證目的"], [
                ["公開 Snapshot", "一個網址、四題觀測、有限證據", "免費", "內容獲客與問題辨識"],
                ["30 日 Decision Loop", "一個 Project、兩次 review、read only", "NT$1,980 至 2,980", "測付款與第二次使用"],
                ["Practitioner Workspace", "3 個 Projects、月度 review、匯出", "NT$2,490 至 3,990 每月", "測小顧問續訂"],
                ["Studio Workspace", "10 個 Projects、品牌化報告與協作", "NT$5,990 至 8,990 每月", "第二階段測多客戶價值"],
            ]),
            ("p", "H", "以上價格是預售實驗範圍，不是市場成交價。應在測試開始前凍結方案與範圍，避免為了成交逐案加入大量人工服務。"),
            ("kpis", [
                ("40", "小方案客戶", "若 ARPA 為 NT$2,490，約 NT$99,600 MRR"),
                ("17", "Studio 客戶", "若 ARPA 為 NT$5,990，約 NT$101,830 MRR"),
                ("≤15%", "Provider 成本", "占訂閱營收的建議上限"),
                ("≤20 分", "人工時間", "成熟後每 Project 每月"),
            ]),
            ("p", "I", "上述算式只顯示成為小型產品所需的客戶量，不是需求預測。若取得 40 至 50 個付費團隊仍高度依賴逐案銷售或人工分析，模式就不是可擴張 SaaS。"),
        ],
    },
    {
        "id": "distribution",
        "eyebrow": "09 Distribution",
        "title": "無人脈與弱 BD 條件下的獲客設計",
        "blocks": [
            ("p", "R", "把自媒體視為產品的分發系統，而不是單純曝光。每篇內容都應讓目標顧問完成一個可觀察行為，例如提交網站、查看來源證據、選擇要追查的問題，或設定下一輪回看。"),
            ("table", ["內容資產", "受眾問題", "產品入口", "成功訊號"], [
                ["餐飲 AI 搜尋案例", "為何某品牌被推薦或忽略", "查看 evidence 與提交自己的網址", "非熟人完成提交"],
                ["每月變化拆解", "這個月哪些變化值得做", "選一個 Agent 問題", "查看來源或建立回查"],
                ["餐飲資料錯誤案例", "品牌與分店資訊不一致", "下載檢查表或建立 Project", "匯入實際資料"],
                ["顧問交付模板", "如何向客戶說明 AI visibility", "產生可分享 decision brief", "第二次回訪或付費"],
            ]),
            ("p", "I", "內容通路適合創辦人現況，但速度慢，且有吸引免費好奇流量的風險。必須分開記錄內容曝光、合格落地頁工作階段、網址提交、證據查看、第二次使用與付款，不能用按讚或追蹤數代替需求。"),
            ("callout", "R", "分發門檻", "在至少 150 個符合目標內容的落地頁 sessions 中，取得 8 個非熟人提交，其中至少 3 個查看證據或設定下一輪回看。流量不足只代表分發未驗證；有合格流量卻無行為，才是需求警訊。"),
        ],
    },
    {
        "id": "competition",
        "eyebrow": "10 Competitive position",
        "title": "競爭策略與不對稱選擇",
        "blocks": [
            ("table", ["競爭者", "主要優勢", "GeoCheck 不應跟隨", "可攻擊位置"], [
                ["Hogiah", "繁中、在地案例、LINE、內容與餐飲定位", "引擎數、內容量與在地學院規模", "顧問多客戶交付與證據決策紀錄"],
                ["Ahrefs 與 Semrush", "資料、品牌、既有工作流與 bundle", "通用 SEO Suite", "台灣餐飲 entity quality 與低摩擦月報"],
                ["Peec 與 Otterly", "成熟監測、agency 功能、prompt 工作流", "更完整的通用監測功能", "保守證據規則與餐飲流程模板"],
                ["Google GSC", "第一方免費曝光資料", "單一平台數據顯示", "跨引擎觀測與人工決策層"],
                ["試算表與 ChatGPT", "便宜、彈性高", "通用聊天", "版本、證據、重跑與客戶交付一致性"],
            ]),
            ("p", "F", "Ahrefs 已公開 demand adjusted volume，Otterly 已有 GSC 到 prompt 的工作流，因此 Demand weighting 或 GSC prompt generation 不能單獨當作新品類。[S06][S07][S09]"),
            ("p", "I", "GeoCheck 應把差異化放在工作完成品質：能否正確處理資料不足、找到品牌或分店混淆、引用可核對的 evidence，並把一次建議變成下次可驗證的決策。"),
        ],
    },
    {
        "id": "roadmap",
        "eyebrow": "11 Product roadmap",
        "title": "六個月開發與驗證 Roadmap",
        "blocks": [
            ("roadmap", [
                ("M0", "Evidence Agent 評測", "建立 20 個 fixture，涵蓋有效提及、官方引用、provider 失敗、少於兩個有效 run、跨次變化與相似品牌。", "20 案不把 unknown 當零、不杜撰來源；至少 18 案符合人工預期。"),
                ("M1", "公開案例與單題預覽", "發布三個可查案例，讓使用者提交網址並選一個 bounded question。", "150 個合格 sessions、8 個非熟人提交、3 個 evidence 或回查行為。"),
                ("M2", "30 日付費 Decision Loop", "最多三個 Project，兩次 review，控制每案人工時間與 provider 成本。", "至少 2 個付費；或 1 個付費加 2 個資料授權者完成第二輪。"),
                ("M3", "最小 Workspace", "只在 M2 通過後建立 Project、Evidence History、匯入與 Decision Ledger。", "資料格式可重用；顧問確實回看同一 Project。"),
                ("M4", "餐飲 GeoBooks", "將已被人完成並確認的月報、來源缺口與分店混淆流程產品化。", "至少兩個顧問重跑同一 workflow，且不需逐案改提示。"),
                ("M5", "有限 Monitoring 與 GSC", "先支援 CSV；只有重複需求才做 OAuth。加入週或月 sentinel runs 與變更提醒。", "低樣本可 abstain；provider 成本不超過收入 15%。"),
                ("M6", "商業 Gate", "檢查續訂、交付工時、毛利、合格流量與獲客可重複性。", "通過則擴大窄 SaaS；一次性價值改走產品化服務；無付款或無第二輪則停止。"),
            ]),
        ],
    },
    {
        "id": "agent-data",
        "eyebrow": "12 Agent and data",
        "title": "Agent 素材與資料治理",
        "blocks": [
            ("p", "R", "第一階段不要訓練或微調模型。先建立版本化的餐飲 context package、固定工具、輸出 schema 與評測 fixtures。只有累積反覆錯誤模式後，才判斷是否需要訓練。"),
            ("table", ["素材", "第一版內容", "品質要求"], [
                ["Entity schema", "品牌、分店、官方 URL、地址、菜單、價位、訂位、外送與時效", "每個欄位有來源與更新時間"],
                ["Prompt library", "地點、餐種、場合、價位、飲食需求與轉換意圖", "記錄來源、版本與適用條件"],
                ["Evidence rules", "有效答案、官方來源、品牌分店消歧、unknown 與 parser 限制", "任何事實可回到 evidence ID"],
                ["Action taxonomy", "檢查菜單頁、分店頁、訂位資訊、結構化資料與來源缺口", "建議不等於執行或效果"],
                ["GeoBooks", "解釋變化、來源消失、分店混淆與月度 review", "先由人完成並審核，再允許重跑"],
                ["Fixtures and evals", "正常、模糊、未知、加盟、多分店、價格改動與 provider 失敗", "保存預期答案與失敗類型"],
            ]),
            ("callout", "R", "資料界線", "公開方法、公開案例、客戶私有資料與專案研究資料必須分開。客戶資料不得因為可形成護城河就被未經同意地拿去訓練或跨客戶使用。"),
        ],
    },
    {
        "id": "metrics",
        "eyebrow": "13 Metrics and economics",
        "title": "商業指標與單位經濟",
        "blocks": [
            ("table", ["層級", "主要指標", "為何重要", "不接受的替代指標"], [
                ["分發", "合格 sessions 到網址提交率", "內容是否帶來目標使用者行為", "瀏覽量、按讚、追蹤數"],
                ["啟用", "查看 evidence 或完成第一個決策", "是否超過一次性分數好奇", "註冊帳號"],
                ["價值", "每 Project 至少一項被採用決策", "是否幫顧問完成工作", "Agent 對話次數"],
                ["留存", "30 日內第二次 review 同一 Project", "訂閱理由是否存在", "每日登入"],
                ["效率", "每份決策包人工分鐘數", "一至兩人能否交付", "自動化功能數"],
                ["經濟", "ARPA、provider 成本率、退款與直接毛利", "是否能產品化", "名義 MRR 不扣成本"],
            ]),
            ("p", "H", "進入擴張階段的建議條件是：付費轉換存在、至少一半付費試點完成第二輪、成熟後每 Project 每月人工低於 20 分鐘、provider 成本低於收入 15%、直接毛利高於 70%。門檻是投資紀律，不是已證實 benchmark。"),
        ],
    },
    {
        "id": "gates",
        "eyebrow": "14 Validation gates",
        "title": "三個不能跳過的驗證實驗",
        "blocks": [
            ("table", ["實驗", "承諾等級", "成功條件", "失敗後決策"], [
                ["Evidence Agent 正確性", "內部可重現結果", "20 案無杜撰與 unknown 錯誤，至少 18 案符合預期", "縮小工具與 schema，不加聊天功能"],
                ["公開案例行為", "非熟人提交並查看證據", "8 個提交，3 個 evidence 或回查行為", "若流量足夠仍無行為，改 TA 或問題"],
                ["30 日預售", "付款或資料授權加第二輪", "2 個付費，或 1 個付費加 2 個第二輪資料承諾", "停止 Agent 訂閱假設，保留研究或一次性服務"],
            ]),
            ("callout", "R", "不可用功能補失敗", "若實驗失敗，先判斷是流量、目標客群、問題、價格或交付失敗。不能用新增引擎、更多 Dashboard 或更自由的 Agent 延後停損。"),
        ],
    },
    {
        "id": "risks",
        "eyebrow": "15 Risk register",
        "title": "主要風險與降低方式",
        "blocks": [
            ("table", ["風險", "可能性", "影響", "早期訊號", "處置"], [
                ["顧問市場太小", "高", "高", "有內容流量但合格提交少", "放寬至多據點在地品牌，不立即擴全產業"],
                ["付費只買執行", "高", "高", "喜歡報告但不願訂閱", "改為顧問交付工具或產品化服務"],
                ["Google 與 Suite 商品化", "高", "中高", "使用者說現有工具已足夠", "聚焦跨來源證據、工作流與決策歷史"],
                ["Agent 產生錯誤建議", "中高", "高", "fixture 失敗、引用對不上", "限制工具、schema、信心標示與人工批准"],
                ["API 成本與不穩定", "中", "中高", "重試多、成本超過收入 15%", "固定少量 sentinel、快取、預算上限、abstain"],
                ["創辦人分發不足", "高", "高", "無法取得 150 個合格 sessions", "先改善內容渠道；不把需求判為失敗"],
                ["餐飲資料變動快", "中高", "中", "菜單、價格、分店資訊頻繁失效", "每欄位保存來源與更新時間"],
            ]),
        ],
    },
    {
        "id": "not-build",
        "eyebrow": "16 Scope discipline",
        "title": "未來六個月絕對不要做",
        "blocks": [
            ("bullets", "R", [
                "多 Agent 協作平台、任意工具市場或通用 Growth OS。",
                "自動產文、自動發布、自動修改客戶網站或無人工批准的外部動作。",
                "一次支援所有產業、所有國家、所有 AI 引擎與每日大量 prompt 掃描。",
                "Enterprise API、SSO、複雜權限、白標客製與長採購流程。",
                "把 GSC demand、AI visibility、GA4 referral 與 conversion 混成一個未驗證的 Reality Score。",
                "在付款與第二次使用成立前，重做完整 onboarding、billing、通知中心或精緻 Dashboard。",
            ]),
        ],
    },
    {
        "id": "kill",
        "eyebrow": "17 Adversarial plan",
        "title": "競爭對手如何在六個月內擊敗 GeoCheck",
        "blocks": [
            ("p", "I", "最簡單的攻擊不是複製演算法，而是把 GeoCheck 的主要價值免費 bundle 到既有工具與服務。"),
            ("table", ["攻擊者", "六個月行動", "GeoCheck 受到的傷害", "必要防禦"], [
                ["Hogiah", "推出多客戶 workspace、餐飲報告與顧問推薦分潤", "在地定位與通路被包圍", "先建立顧問決策流程與可轉移歷史"],
                ["SEO Suite", "將 GSC AI 報告、Agent 問答與有限 monitoring 納入現有方案", "Dashboard 與需求加權失去收費空間", "餐飲 entity、證據品質與交付模板"],
                ["在地 Agency", "免費健檢搭配內容、評論、廣告代操", "買家更願意為執行付費", "只服務需要保留判斷主導權的顧問"],
                ["開源團隊", "提供免費 tracker 與 ChatGPT 報表 prompt", "價格被壓低", "版本、回查、資料治理與低工時"],
            ]),
            ("callout", "R", "六個月防禦目標", "不要追求競品無法複製的功能。先建立競品複製後仍缺少的使用歷史：某位顧問對某個品牌做過什麼決策、依據是什麼、下一輪結果如何。若使用者不願留下這種歷史，護城河假設失敗。"),
        ],
    },
    {
        "id": "ask",
        "eyebrow": "18 Investment use",
        "title": "建議投入方式與六個月決策",
        "blocks": [
            ("p", "R", "本案目前適合的是里程碑式投入，而不是一次投入完整 SaaS 建設。資源應先買到需求、付款與留存證據，再增加工程複雜度。"),
            ("table", ["資源", "建議用途", "上限或條件"], [
                ["創辦人時間", "公開案例、三個付費 Project、決策流程拆解", "同時最多三個 Project"],
                ["模型與 provider", "fixture、少量 sentinel runs、兩次 review", "各實驗預設硬上限並記錄重試成本"],
                ["工程", "Evidence schema、Project history、Decision Ledger", "只有上一階段行為門檻通過才開發"],
                ["內容", "三個餐飲案例、方法拆解、顧問模板", "每篇都有產品行為 CTA"],
                ["外部協助", "一名合格顧問審閱與資料權限諮詢", "優先取得批評，不用品牌背書取代驗證"],
            ]),
            ("callout", "R", "六個月決策規則", "Continue：付款、第二輪、工時與毛利同時通過。Modify：有人付費但主要價值在一次性診斷或人工交付。Pivot：顧問不付費，但多據點品牌顯示明確持續需求。Stop：合格流量與多次提案後仍沒有付款或第二輪行為。"),
            ("p", "U", "本企劃沒有估算募資金額、台灣 TAM、CAC、LTV 或現金跑道，因為目前缺乏成交、留存、可接觸 TA 數量與實際 provider 成本資料。填入漂亮數字會降低投資判斷品質。"),
        ],
    },
]


def status_badge(code: str) -> str:
    label, css = LABELS[code]
    return f'<span class="badge {css}">{label}</span>'


def refs_to_links(text: str) -> str:
    safe = escape(text)
    return re.sub(r"\[(S\d{2})\]", lambda m: f'<a class="cite" href="#{m.group(1).lower()}">[{m.group(1)}]</a>', safe)


def render_block(block) -> str:
    kind = block[0]
    if kind == "p":
        _, status, text = block
        return f'<p class="claim" data-kind="{LABELS[status][1]}">{status_badge(status)} {refs_to_links(text)}</p>'
    if kind == "callout":
        _, status, title, text = block
        return f'<aside class="callout {LABELS[status][1]}" data-kind="{LABELS[status][1]}"><div>{status_badge(status)}<h3>{escape(title)}</h3></div><p>{refs_to_links(text)}</p></aside>'
    if kind == "table":
        _, headers, rows = block
        head = "".join(f"<th>{escape(h)}</th>" for h in headers)
        body = "".join("<tr>" + "".join(f"<td>{refs_to_links(str(c))}</td>" for c in row) + "</tr>" for row in rows)
        return f'<div class="table-wrap"><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>'
    if kind == "kpis":
        cards = "".join(f'<div class="metric"><strong>{escape(v)}</strong><span>{escape(label)}</span><small>{escape(note)}</small></div>' for v, label, note in block[1])
        return f'<div class="metrics">{cards}</div>'
    if kind == "flow":
        steps = "".join(f'<div class="flow-step"><span>{i:02d}</span><h3>{escape(title)}</h3><b>{escape(sub)}</b><p>{escape(note)}</p></div>' for i, (title, sub, note) in enumerate(block[1], 1))
        return f'<div class="flow">{steps}</div>'
    if kind == "roadmap":
        items = "".join(f'<article class="road-item"><div class="month">{escape(m)}</div><div><h3>{escape(title)}</h3><p>{escape(work)}</p><p class="gate"><b>進入門檻</b> {escape(gate)}</p></div></article>' for m, title, work, gate in block[1])
        return f'<div class="roadmap">{items}</div>'
    if kind == "bullets":
        _, status, items = block
        lis = "".join(f"<li>{refs_to_links(item)}</li>" for item in items)
        return f'<div class="bullet-block" data-kind="{LABELS[status][1]}">{status_badge(status)}<ul>{lis}</ul></div>'
    raise ValueError(kind)


def build_html() -> None:
    nav = "".join(f'<a href="#{s["id"]}"><span>{i:02d}</span>{escape(s["title"])}</a>' for i, s in enumerate(SECTIONS, 1))
    sections = []
    for sec in SECTIONS:
        blocks = "".join(render_block(b) for b in sec["blocks"])
        sections.append(f'<section id="{sec["id"]}"><div class="section-head"><p>{escape(sec["eyebrow"])}</p><h2>{escape(sec["title"])}</h2></div>{blocks}</section>')
    source_html = "".join(
        f'<li id="{sid.lower()}"><span>{sid}</span><div><strong>{escape(title)}</strong><small>{escape(kind)}</small><a href="{escape(url)}">{escape(url)}</a><p>{escape(note)}</p></div></li>'
        for sid, title, kind, url, note in SOURCES
    )
    html = f'''<!doctype html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="description" content="GeoCheck 投資人版本商業企劃書，涵蓋市場、客群、商業模式、Agent Workspace 與六個月 Roadmap。">
<title>GeoCheck 商業企劃書</title>
<style>
:root{{--paper:#f3efe6;--ink:#17221f;--muted:#69716d;--line:#c8c4b9;--green:#123d35;--mint:#c6e0d1;--orange:#e6673c;--cream:#fffdf7;--yellow:#f2d477;--red:#9b3b2c}}
*{{box-sizing:border-box}}html{{scroll-behavior:smooth}}body{{margin:0;background:var(--paper);color:var(--ink);font-family:"Noto Sans TC","Microsoft JhengHei",sans-serif;line-height:1.75}}
body:before{{content:"";position:fixed;inset:0;pointer-events:none;opacity:.18;background-image:radial-gradient(#17221f 0.45px,transparent 0.45px);background-size:6px 6px;z-index:20}}
a{{color:inherit}}.shell{{display:grid;grid-template-columns:270px minmax(0,1fr)}}
.sidebar{{position:sticky;top:0;height:100vh;padding:28px 20px;background:var(--green);color:#fff;overflow:auto;z-index:30}}.brand{{display:flex;align-items:center;gap:10px;font-family:Georgia,"Noto Serif TC",serif;font-size:24px;margin-bottom:25px}}.brand i{{width:13px;height:13px;border-radius:50%;background:var(--orange);box-shadow:18px 0 0 var(--yellow)}}
.sidebar nav{{display:grid;gap:2px}}.sidebar nav a{{display:grid;grid-template-columns:30px 1fr;gap:4px;text-decoration:none;padding:7px 8px;border-radius:8px;font-size:12px;color:#c9d9d2}}.sidebar nav a:hover{{background:#ffffff14;color:#fff}}.sidebar nav span{{font-family:ui-monospace,monospace;color:#82a99a}}
.legend{{border-top:1px solid #ffffff28;margin-top:22px;padding-top:18px;font-size:11px;color:#c9d9d2}}.legend button{{border:1px solid #ffffff35;background:transparent;color:#fff;border-radius:20px;padding:5px 8px;margin:2px;cursor:pointer}}.legend button.active{{background:var(--orange);border-color:var(--orange)}}
.main{{min-width:0}}.hero{{min-height:88vh;display:grid;grid-template-columns:1.35fr .65fr;padding:7vw 7vw 5vw;gap:48px;align-items:end;border-bottom:1px solid var(--line);position:relative;overflow:hidden}}.hero:after{{content:"MODIFY";position:absolute;right:-20px;top:28px;font:900 12vw/1 Georgia,serif;color:#123d350d;transform:rotate(-6deg)}}
.kicker{{text-transform:uppercase;letter-spacing:.18em;font-size:12px;font-weight:800;color:var(--orange)}}h1,h2,h3{{font-family:Georgia,"Noto Serif TC","Microsoft JhengHei",serif}}h1{{font-size:clamp(54px,8vw,112px);line-height:.92;letter-spacing:-.06em;margin:20px 0 28px;max-width:900px}}.hero .deck{{font-size:clamp(18px,2vw,27px);line-height:1.45;max-width:770px}}.hero-meta{{border-left:4px solid var(--orange);padding-left:20px}}.hero-meta strong{{display:block;font:700 44px/1 Georgia,serif;color:var(--green)}}.hero-meta span{{display:block;color:var(--muted);margin:8px 0 28px}}.hero-meta p{{font-size:13px}}
main article{{max-width:1120px;margin:auto;padding:0 6vw 100px}}section{{padding:84px 0 36px;border-bottom:1px solid var(--line);scroll-margin-top:20px}}.section-head p{{font:700 11px/1 ui-monospace,monospace;text-transform:uppercase;letter-spacing:.16em;color:var(--orange)}}.section-head h2{{font-size:clamp(34px,5vw,58px);line-height:1.08;max-width:820px;margin:14px 0 34px;letter-spacing:-.035em}}
.claim{{font-size:17px;max-width:920px;margin:20px 0}}.badge{{display:inline-block;vertical-align:2px;border-radius:999px;padding:2px 8px;margin-right:7px;font:700 10px/1.7 ui-monospace,monospace;letter-spacing:.08em;background:#ddd}}.badge.fact{{background:#c8e0d4}}.badge.observed{{background:#d9d2ee}}.badge.inference{{background:#e9d49d}}.badge.hypothesis{{background:#f5c7b4}}.badge.unknown{{background:#d3d3cf}}.badge.recommendation{{background:#173f36;color:#fff}}
.callout{{display:grid;grid-template-columns:minmax(190px,.35fr) 1fr;gap:28px;margin:34px 0;padding:28px;border:1px solid var(--ink);background:var(--cream);box-shadow:8px 8px 0 var(--ink)}}.callout h3{{font-size:25px;margin:12px 0 0}}.callout p{{font-size:18px;margin:0}}.callout.hypothesis{{box-shadow:8px 8px 0 var(--orange)}}.callout.unknown{{box-shadow:8px 8px 0 #999}}
.table-wrap{{overflow:auto;margin:30px 0}}table{{width:100%;border-collapse:collapse;background:#fffaf0;font-size:14px}}th{{background:var(--green);color:#fff;text-align:left;padding:13px 14px;vertical-align:top}}td{{border:1px solid #d3cec2;padding:12px 14px;vertical-align:top}}tbody tr:nth-child(even){{background:#eff2eb}}
.metrics{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin:30px 0}}.metric{{border-top:6px solid var(--orange);background:var(--cream);padding:22px 18px;min-height:155px}}.metric strong{{display:block;font:700 38px/1 Georgia,serif}}.metric span{{display:block;font-weight:800;margin-top:13px}}.metric small{{display:block;color:var(--muted);line-height:1.4;margin-top:6px}}
.flow{{display:grid;grid-template-columns:repeat(5,1fr);gap:1px;background:var(--ink);border:1px solid var(--ink);margin:34px 0}}.flow-step{{background:var(--cream);padding:19px;min-height:200px}}.flow-step span{{font:700 12px ui-monospace;color:var(--orange)}}.flow-step h3{{font-size:22px;margin:26px 0 2px}}.flow-step b{{font-size:12px}}.flow-step p{{font-size:12px;color:var(--muted);line-height:1.5}}
.roadmap{{position:relative;margin:32px 0}}.roadmap:before{{content:"";position:absolute;left:41px;top:20px;bottom:30px;width:2px;background:var(--green)}}.road-item{{position:relative;display:grid;grid-template-columns:84px 1fr;gap:24px;margin:0 0 18px}}.month{{position:relative;z-index:1;width:84px;height:54px;display:grid;place-items:center;background:var(--green);color:#fff;font:800 18px ui-monospace,monospace}}.road-item>div:last-child{{background:var(--cream);padding:17px 22px;border:1px solid var(--line)}}.road-item h3{{margin:0;font-size:22px}}.road-item p{{margin:7px 0}}.gate{{color:var(--green);font-size:13px}}
.bullet-block{{display:grid;grid-template-columns:90px 1fr;gap:18px;margin:26px 0}}.bullet-block ul{{margin:0;padding-left:22px}}.bullet-block li{{margin:10px 0;font-size:17px}}
.cite{{font:700 11px ui-monospace;color:var(--red);text-decoration:none}}.sources{{list-style:none;padding:0;display:grid;gap:10px}}.sources li{{display:grid;grid-template-columns:54px 1fr;gap:12px;padding:16px;background:var(--cream);border:1px solid var(--line)}}.sources li>span{{font:800 13px ui-monospace;color:var(--orange)}}.sources strong,.sources small,.sources a{{display:block}}.sources small{{color:var(--muted)}}.sources a{{font-size:11px;color:var(--green);overflow-wrap:anywhere}}.sources p{{font-size:13px;margin:5px 0 0}}
.calculator{{background:var(--green);color:#fff;padding:34px;margin:42px 0;display:grid;grid-template-columns:1fr 1fr;gap:34px}}.calculator h3{{font-size:28px;margin:0 0 10px}}.calculator label{{display:block;margin:12px 0 5px}}.calculator input{{width:100%;accent-color:var(--orange)}}.calc-result{{display:grid;place-items:center;border:1px solid #ffffff55;text-align:center}}.calc-result strong{{font:700 52px/1 Georgia,serif}}.calc-result span{{font-size:12px;color:#cad9d3}}
.filtered{{opacity:.12;filter:grayscale(1)}}footer{{padding:45px 7vw;background:#102c26;color:#cfe0d9;font-size:12px}}@media(max-width:900px){{.shell{{display:block}}.sidebar{{position:relative;height:auto}}.sidebar nav{{grid-template-columns:repeat(2,1fr)}}.hero{{grid-template-columns:1fr;min-height:auto;padding:70px 7vw}}.metrics{{grid-template-columns:repeat(2,1fr)}}.flow{{grid-template-columns:1fr 1fr}}.callout,.calculator{{grid-template-columns:1fr}}}}@media(max-width:560px){{.sidebar nav{{display:none}}h1{{font-size:55px}}.metrics,.flow{{grid-template-columns:1fr}}.road-item{{grid-template-columns:64px 1fr}}.month{{width:64px}}.callout{{box-shadow:5px 5px 0 var(--ink)}}}}
@media print{{.sidebar{{display:none}}.shell{{display:block}}.hero{{min-height:auto;padding:1cm}}main article{{max-width:none;padding:0 1cm}}section{{page-break-before:always;border:0;padding-top:1cm}}.callout{{box-shadow:none}}}}
</style>
</head>
<body>
<div class="shell">
<aside class="sidebar"><div class="brand"><i></i>GeoCheck</div><nav>{nav}<a href="#sources"><span>19</span>來源與證據</a></nav><div class="legend"><p>只顯示特定主張</p><button class="active" data-filter="all">全部</button><button data-filter="fact">事實</button><button data-filter="inference">推論</button><button data-filter="hypothesis">假設</button><button data-filter="unknown">未知</button><button data-filter="recommendation">建議</button></div></aside>
<div class="main">
<header class="hero"><div><p class="kicker">Investor edition · 2026 09 07</p><h1>GeoCheck<br>商業企劃書</h1><p class="deck">把一次性 AI 可見度檢測，改造成服務餐飲顧問的 Evidence Workspace。建議進行六個月條件式驗證，先證明付款與第二次使用，再擴大產品。</p></div><div class="hero-meta"><strong>MODIFY</strong><span>有條件修改方向</span><p>對象：投資人與創辦人<br>範圍：市場、TA、產品、商業模式、分發、Roadmap<br>證據日：2026 年 9 月 7 日</p></div></header>
<main><article>{''.join(sections)}
<section id="calculator"><div class="section-head"><p>Scenario model</p><h2>收入情境試算</h2></div><p class="claim" data-kind="hypothesis">{status_badge('H')}這是理解規模的試算器，不是營收預測。調整平均月費與付費團隊數，查看對應 MRR。</p><div class="calculator"><div><h3>Micro SaaS 情境</h3><label>平均月費 NT$ <b id="price-label">2,990</b></label><input id="price" type="range" min="1000" max="9000" step="250" value="2990"><label>付費團隊數 <b id="customers-label">40</b></label><input id="customers" type="range" min="1" max="150" value="40"></div><div class="calc-result"><div><span>MONTHLY RECURRING REVENUE</span><strong id="mrr">NT$119,600</strong><p id="arr">年化 NT$1,435,200</p></div></div></div></section>
<section id="sources"><div class="section-head"><p>19 Evidence register</p><h2>來源與證據限制</h2></div><p>供應商官網只能證明公開功能、價格或自述；不能證明成交、留存或成效。國際收入資料不能直接外推台灣 TA。所有本地重大方向仍需使用者確認並記錄於 Decision Log。</p><ol class="sources">{source_html}</ol></section>
</article></main><footer>GeoCheck Business Plan · Evidence date 2026 09 07 · Major positioning and commercial decisions remain proposals until recorded in the project decision log.</footer>
</div></div>
<script>
const money=n=>'NT$'+Math.round(n).toLocaleString('zh-TW');
const price=document.querySelector('#price'), customers=document.querySelector('#customers');
function calc(){{const p=+price.value,c=+customers.value,m=p*c;document.querySelector('#price-label').textContent=p.toLocaleString();document.querySelector('#customers-label').textContent=c;document.querySelector('#mrr').textContent=money(m);document.querySelector('#arr').textContent='年化 '+money(m*12)}}
price.addEventListener('input',calc);customers.addEventListener('input',calc);calc();
document.querySelectorAll('[data-filter]').forEach(btn=>btn.addEventListener('click',()=>{{document.querySelectorAll('[data-filter]').forEach(b=>b.classList.remove('active'));btn.classList.add('active');const f=btn.dataset.filter;document.querySelectorAll('[data-kind]').forEach(el=>el.classList.toggle('filtered',f!=='all'&&el.dataset.kind!==f));}}));
</script>
</body></html>'''
    HTML_PATH.write_text(html, encoding="utf-8")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def remove_paragraph_borders(paragraph_or_style) -> None:
    element = paragraph_or_style._element
    p_pr = element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("GeoCheck 商業企劃書   |   ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(105, 113, 109)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def style_run(run, font="Microsoft JhengHei", size=None, color=None, bold=None) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_claim(doc: Document, status: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(f"[{LABELS[status][0]}] ")
    style_run(r, size=10, color="9B3B2C" if status in ("H", "U") else "123D35", bold=True)
    r = p.add_run(text)
    style_run(r, size=10.3, color="17221F")


def add_callout(doc: Document, status: str, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EFF2EB" if status == "R" else "FFF3EA")
    set_cell_border(cell, "B9C8C1", "8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"[{LABELS[status][0]}] {title}")
    style_run(r, size=11.5, color="000000", bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    style_run(r, size=9.8, color="17221F")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers, rows) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for idx, h in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, "123D35")
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(h))
        style_run(r, size=8.5, color="FFFFFF", bold=True)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for idx, value in enumerate(row):
            cell = cells[idx]
            if row_idx % 2:
                set_cell_shading(cell, "F3F6F4")
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            style_run(r, size=8.2, color="17221F")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_flow(doc: Document, steps) -> None:
    table = doc.add_table(rows=1, cols=len(steps))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    prevent_row_split(table.rows[0])
    for idx, (title, sub, note) in enumerate(steps):
        cell = table.cell(0, idx)
        set_cell_shading(cell, "FFFDF7" if idx % 2 == 0 else "EFF2EB")
        set_cell_border(cell, "B9C8C1")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{idx + 1:02d}\n{title}\n")
        style_run(r, size=9.5, color="123D35", bold=True)
        r = p.add_run(f"{sub}\n")
        style_run(r, size=8.2, color="000000", bold=True)
        r = p.add_run(note)
        style_run(r, size=7.5, color="69716D")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_metrics(doc: Document, metrics) -> None:
    table = doc.add_table(rows=1, cols=len(metrics))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    prevent_row_split(table.rows[0])
    for idx, (value, label, note) in enumerate(metrics):
        cell = table.cell(0, idx)
        set_cell_shading(cell, "FFFDF7")
        set_cell_border(cell, "C8C4B9")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{value}\n")
        style_run(r, font="Georgia", size=16, color="123D35", bold=True)
        r = p.add_run(f"{label}\n")
        style_run(r, size=8.3, color="000000", bold=True)
        r = p.add_run(note)
        style_run(r, size=7.2, color="69716D")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullets(doc: Document, status: str, items) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"[{LABELS[status][0]}]")
    style_run(r, size=9, color="123D35", bold=True)
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        style_run(r, size=10, color="17221F")


def add_roadmap(doc: Document, items) -> None:
    for month, title, work, gate in items:
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Cm(2.2)
        table.columns[1].width = Cm(14.6)
        prevent_row_split(table.rows[0])
        left, right = table.rows[0].cells
        left.width = Cm(2.2)
        right.width = Cm(14.6)
        set_cell_shading(left, "123D35")
        set_cell_shading(right, "FFFDF7")
        set_cell_border(left, "123D35")
        set_cell_border(right, "C8C4B9")
        p = left.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(month)
        style_run(r, size=12, color="FFFFFF", bold=True)
        p = right.paragraphs[0]
        r = p.add_run(title + "\n")
        style_run(r, size=10.5, color="000000", bold=True)
        r = p.add_run(work + "\n")
        style_run(r, size=8.7, color="17221F")
        r = p.add_run("進入門檻  " + gate)
        style_run(r, size=8, color="123D35", bold=True)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(1)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    keep_with_next(p)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft JhengHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    styles["Normal"].font.size = Pt(10.3)
    styles["Normal"].font.color.rgb = RGBColor.from_string("17221F")
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.22
    for name, size in (("Title", 28), ("Heading 1", 20), ("Heading 2", 13)):
        style = styles[name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(8)
        remove_paragraph_borders(style)

    title = doc.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(70)
    title.paragraph_format.space_after = Pt(16)
    r = title.add_run("GeoCheck 商業企劃書")
    style_run(r, size=30, color="000000", bold=True)
    p = doc.add_paragraph()
    r = p.add_run("投資人版本")
    style_run(r, size=15, color="000000", bold=True)
    p.paragraph_format.space_after = Pt(34)
    p = doc.add_paragraph()
    r = p.add_run("建議決策  MODIFY")
    style_run(r, font="Georgia", size=20, color="123D35", bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("把一次性 AI 可見度檢測改造成服務餐飲顧問的 Evidence Workspace。先進行六個月條件式驗證，證明付款與第二次使用後，再投資完整 SaaS。")
    style_run(r, size=12, color="17221F")
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row, vals in zip(meta.rows, [("文件日期", "2026 年 9 月 7 日"), ("讀者", "潛在投資人與創辦人"), ("決策狀態", "研究建議，尚未寫入正式 Decision Log"), ("證據範圍", "本地產品文件與截至文件日期的公開資料")]):
        for i, value in enumerate(vals):
            cell = row.cells[i]
            set_cell_border(cell, "D9D9D9")
            if i == 0:
                set_cell_shading(cell, "EFF2EB")
            r = cell.paragraphs[0].add_run(value)
            style_run(r, size=9, color="000000", bold=(i == 0))
    doc.add_page_break()

    add_heading(doc, "投資摘要", 1)
    add_callout(doc, "R", "目前最合理的投資方式", "以六個月、三道驗證門檻投入。第一筆資源用於 Evidence Agent 正確性、公開案例與三個付費 Project；只有付款與第二次使用成立，才建 Project Workspace、Monitoring 與 GSC 整合。")
    add_claim(doc, "F", "國際 AI visibility 類別已有收入與付費方案證據，但台灣小型顧問的付費意願、留存與可接觸規模仍未知。")
    add_claim(doc, "I", "半 Agent 加 Dashboard 比純 Dashboard 更接近顧問的工作，但 Agent 介面本身容易被競品複製。可累積價值來自證據品質、餐飲實體資料、決策歷史與重跑流程。")
    add_claim(doc, "H", "第一 TA 是有 3 至 15 個月約客戶的獨立 SEO 或內容顧問與小型行銷團隊。餐飲品牌是 Project，單店餐廳老闆主要作免費入口。")
    add_heading(doc, "閱讀方式", 2)
    add_claim(doc, "F", "本文用事實、觀察、推論、假設、未知與建議六種標籤。供應商自述不等於成效驗證，未知也不等於零。")
    add_heading(doc, "章節", 2)
    toc_items = [sec["title"] for sec in SECTIONS] + ["收入情境", "來源與證據限制"]
    toc = doc.add_table(rows=10, cols=2)
    toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    toc.autofit = True
    for i, item in enumerate(toc_items, 1):
        row_idx = (i - 1) % 10
        col_idx = (i - 1) // 10
        cell = toc.cell(row_idx, col_idx)
        set_cell_border(cell, "FFFFFF", "0")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{i:02d}  {item}")
        style_run(r, size=8.7, color="17221F")

    for sec_idx, sec in enumerate(SECTIONS, 1):
        p = doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(sec["eyebrow"].upper())
        style_run(r, font="Consolas", size=8, color="E6673C", bold=True)
        add_heading(doc, sec["title"], 1)
        for block in sec["blocks"]:
            kind = block[0]
            if kind == "p":
                add_claim(doc, block[1], block[2])
            elif kind == "callout":
                add_callout(doc, block[1], block[2], block[3])
            elif kind == "table":
                add_table(doc, block[1], block[2])
            elif kind == "kpis":
                add_metrics(doc, block[1])
            elif kind == "flow":
                add_flow(doc, block[1])
            elif kind == "roadmap":
                add_roadmap(doc, block[1])
            elif kind == "bullets":
                add_bullets(doc, block[1], block[2])

    p = doc.add_heading("收入情境", level=1)
    p.paragraph_format.page_break_before = True
    add_claim(doc, "H", "以下數字用來理解規模，不是營收預測。")
    add_table(doc, ["情境", "平均月費", "付費團隊", "MRR", "年化收入"], [
        ["保守", "NT$2,490", "20", "NT$49,800", "NT$597,600"],
        ["基準", "NT$2,990", "40", "NT$119,600", "NT$1,435,200"],
        ["Studio mix", "NT$4,500", "50", "NT$225,000", "NT$2,700,000"],
        ["擴張前門檻", "NT$5,000", "80", "NT$400,000", "NT$4,800,000"],
    ])
    add_claim(doc, "I", "若達到上述客戶數仍需要逐案客製、創辦人親自分析或高成本 API 掃描，這些收入不能被視為可擴張 SaaS。")

    p = doc.add_heading("來源與證據限制", level=1)
    p.paragraph_format.page_break_before = True
    add_claim(doc, "F", "公開價格與功能只能證明供應商如此呈現，不能證明成交、使用頻率或留存。Peec 的 ARR 為公司自報，沒有獨立審計。台灣公開月費方案可證明持續服務存在，不能證明 GeoCheck 的工具型訂閱需求。")
    add_table(doc, ["編號", "來源", "類型", "位置或網址", "用途與限制"], SOURCES)

    for section in doc.sections:
        footer = section.footer
        footer.distance = Cm(0.8)
        add_page_number(footer.paragraphs[0])
    props = doc.core_properties
    props.title = "GeoCheck 商業企劃書"
    props.subject = "AI Visibility Evidence Workspace 商業模式與六個月 Roadmap"
    props.author = "GeoCheck"
    props.keywords = "GeoCheck, GEO, AI Visibility, Agent Workspace, 商業企劃書"
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_html()
    build_docx()
    print(HTML_PATH)
    print(DOCX_PATH)
