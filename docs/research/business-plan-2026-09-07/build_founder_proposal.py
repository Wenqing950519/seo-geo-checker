from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent / "GeoCheck創業提案書_投資人暨U-start版.docx"

GREEN = "174A3A"
LIGHT = "EEF4F1"
GRAY = "5D6762"
PALE = "F5F6F5"
RED = "9B2C2C"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header(row):
    repeat_header(row)


def page_break_before(paragraph):
    paragraph.paragraph_format.page_break_before = True


doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Mm(20)
section.bottom_margin = Mm(18)
section.left_margin = Mm(22)
section.right_margin = Mm(20)
section.header_distance = Mm(8)
section.footer_distance = Mm(8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft JhengHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (
    ("Title", 28, GREEN, 0, 12),
    ("Subtitle", 13, GRAY, 0, 8),
    ("Heading 1", 18, GREEN, 16, 8),
    ("Heading 2", 13, GREEN, 12, 5),
    ("Heading 3", 11, GRAY, 9, 4),
):
    s = styles[name]
    s.font.name = "Microsoft JhengHei"
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.font.bold = name != "Subtitle"
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    ppr = s._element.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is not None:
        ppr.remove(borders)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("GeoCheck 創業提案書｜")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


add_page_number(section.footer.paragraphs[0])


def p(text="", bold_prefix=None, italic=False, align=None):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r = para.add_run(bold_prefix)
        r.bold = True
        para.add_run(text[len(bold_prefix):])
    else:
        r = para.add_run(text)
        r.italic = italic
    return para


def bullet(text, level=0):
    para = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    para.add_run(text)
    return para


def number(text):
    para = doc.add_paragraph(style="List Number")
    para.add_run(text)
    return para


def label_box(label, text, color=LIGHT):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Cm(3.0)
    t.columns[1].width = Cm(13.2)
    row = t.rows[0]
    for c in row.cells:
        set_cell_margins(c, 120, 130, 120, 130)
        set_cell_shading(c, color)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    row.cells[0].paragraphs[0].add_run(label).bold = True
    row.cells[1].paragraphs[0].add_run(text)
    keep_row(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def table(headers, rows, widths=None, font_size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = widths is None
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, GREEN)
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(font_size)
        if widths:
            cell.width = Cm(widths[i])
    keep_row(hdr)
    for row_data in rows:
        row = t.add_row()
        keep_row(row)
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            set_cell_margins(cell)
            if widths:
                cell.width = Cm(widths[i])
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(2)
            run = cell.paragraphs[0].add_run(str(value))
            run.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def h1(text, new_page=False):
    para = doc.add_paragraph(text, style="Heading 1")
    if new_page:
        page_break_before(para)
    return para


def h2(text):
    return doc.add_paragraph(text, style="Heading 2")


def h3(text):
    return doc.add_paragraph(text, style="Heading 3")


# Cover
doc.add_paragraph("GeoCheck", style="Title")
doc.add_paragraph("創業提案書", style="Title")
doc.add_paragraph("投資人暨 U-start 評審閱讀版", style="Subtitle")
p("版本日期：2026 年 9 月 7 日")
p("提案類型：創新服務／AI SaaS／GEO 可觀測量測")
p("文件用途：天使投資初談、學校育成面談、創業競賽書面審查")
doc.add_paragraph()
label_box("一句話說明", "GeoCheck 協助品牌與行銷實作者，看清楚 AI 搜尋是否採用品牌答案、是否引用官方網站，以及下一輪該驗證什麼。", LIGHT)
label_box("目前階段", "已有可運作的單站檢測、正式站供應商驗證、AI 信任值與可追溯報告基礎；尚未完成目標客群、付費意願與重複使用驗證。", PALE)
label_box("本案建議", "先把一次性 Dashboard 發展為 Evidence Workspace：Dashboard 保存證據，受限 Geo Agent 協助解讀、決策與回查；以六個月里程碑證明需求後再擴大。", LIGHT)
doc.add_paragraph()
p("文件誠信聲明：文內將「已驗證事實」「觀察結果」「推論」「建議」「假設」與「未知」分開。所有尚未發生的營收、定價、留存與市場份額均不寫成既有成果。", italic=True)


h1("目錄與閱讀方式", new_page=True)
p("本提案先回答『這是什麼』與『做到哪裡』，再回答『值得投資什麼』。投資人可先讀第 1、4、9、12、15 章；U-start 評審可再讀第 10、11、13、14 章。")
table(["章次", "內容", "主要回答"], [
    ("1–3", "專案與產品", "問題是什麼、GeoCheck 怎麼運作、使用者得到什麼"),
    ("4–5", "現況與技術", "完成了什麼、哪些能力已有證據、哪些仍未知"),
    ("6–8", "市場、客群、競爭", "誰可能買、為什麼現在、憑什麼不被取代"),
    ("9–12", "商業與 Roadmap", "怎麼收費、如何獲客、六個月做什麼、如何停損"),
    ("13–15", "團隊、財務、投資判斷", "執行缺口、資金用途、投資條件與主要風險"),
    ("16–18", "U-start 對照與證據", "評選項目覆蓋、申請資格差距、來源與名詞"),
], widths=[1.8, 4.2, 10.2])
label_box("建議決策", "MODIFY：保留 GeoCheck 已建立的研究與證據基礎，調整為半 Agent＋Dashboard 的窄工作流程；先驗證付款與第二次使用，不直接投入完整通用 SaaS。", "FFF3E8")


h1("1. 專案是什麼", new_page=True)
h2("1.1 背景問題")
p("生成式搜尋逐漸成為消費者研究品牌、產品與服務的入口。品牌開始在意 ChatGPT、Perplexity、Google AI 等系統是否提到自己、引用哪些來源，以及競爭者為何出現在答案裡。現有工具多能提供監測圖表，但圖表不會自動回答：這次變化是否可信、需要處理什麼、處理後如何驗證。")
p("GeoCheck 從一個一次性網站健檢工具出發，嘗試建立可觀測、可回查的 GEO 證據。它不把 AI 的內部想法包裝成可測量事實，也不承諾提高流量或營收；它只在指定查詢、時間、模型與資料範圍內，保存可被檢查的答案採用與引用證據。[S1–S4]")
h2("1.2 一般人如何理解")
p("如果傳統 SEO 工具回答『網站在 Google 排名如何』，GeoCheck 想回答的是：『當使用者用自然語言詢問 AI 時，品牌有沒有進入答案？AI 是否引用品牌的官方來源？我們對這個觀察有多少有效證據？』")
label_box("不做的承諾", "GeoCheck 不保證任何 AI 引擎一定推薦、引用或帶來客戶；目前也不替客戶自動寫文、發文或修改網站。", "FDECEC")
h2("1.3 專案的雙重價值")
table(["價值層", "目前用途", "未來可能形成的產品"], [
    ("研究與可信度", "建立可複現的量測方法、資料治理與白皮書", "可審計的 GEO measurement layer"),
    ("商業與工作流程", "單站檢測與可讀報告", "Evidence Workspace：監測、解讀、決策、驗證"),
], widths=[3.0, 6.2, 7.0])


h1("2. GeoCheck 如何運作", new_page=True)
h2("2.1 現行單站檢測流程")
number("使用者輸入品牌網站，或補充自己想觀察的非品牌問題。")
number("系統讀取網站內容，辨識產業與品牌實體，產生候選查詢。")
number("後端選出四個不重複的 discovery queries，交由 Perplexity 取得可觀測答案與引用。")
number("系統檢查品牌是否被採用、引用網址是否屬於已驗證的第一方官方來源。")
number("報告呈現 AI 信任值、原始證據、站內準備度與優先處理建議；資料不足時顯示 unknown，而不是補成 0。[S2–S4]")
h2("2.2 AI 信任值的定位")
p("目前產品指標 AI Trust Index／AI 信任值 v1.0.0，由答案採用率 65% 與已驗證第一方 URL 來源證據率 35% 組成。這是產品穩定性規則，用來一致呈現觀測結果，不是學術界公認常數，也不是全市場曝光率。[S3]")
table(["構成", "回答的問題", "不能推論"], [
    ("答案採用率 65%", "有效回答中，AI 是否採用該品牌", "不能代表所有 AI 使用者或所有問題"),
    ("來源證據率 35%", "有效回答是否引用已驗證官方 URL", "不能證明引用內容正確或造成購買"),
    ("unknown 規則", "證據不足時停止給分", "不能把未知解讀為品牌表現為零"),
], widths=[3.3, 6.4, 6.5])
h2("2.3 半 Agent＋Dashboard 的下一版")
p("Dashboard 保留每次觀測、證據、分母、版本與變化；Geo Agent 只能讀取這些已保存資料，產出附 evidence ID 的解讀卡。使用者可以接受、拒絕、延後或要求人工確認，決策會寫入 Decision Ledger，下一輪再檢查是否出現預期變化。這個閉環稱為 Evidence → Diagnose → Decide → Verify。[S6]")


h1("3. 使用者為什麼可能需要它", new_page=True)
h2("3.1 候選工作情境")
p("一名 SEO／內容顧問每月要向多個客戶交代：AI 搜尋裡發生了什麼、是正常波動還是值得處理、下一個月要追什麼。現在他通常在監測工具、GA4、GSC、試算表與聊天模型間搬資料，再人工寫報告。GeoCheck 的價值假設，是把證據與判斷放在同一個可回查的 Project 中，減少拼接時間與無法追溯的結論。")
h2("3.2 建議的首批使用者")
table(["角色", "使用觸發", "可能付費原因", "目前證據"], [
    ("獨立 SEO／內容顧問", "月報、客戶追問、重要頁面改版", "縮短交付時間、提高解釋可信度", "假設；尚無訪談或付款"),
    ("小型行銷工作室", "同時管理 3–15 個長期客戶", "多 Project 與可重複報告流程", "假設；競品有 agency 方案"),
    ("In-house SEO／Growth generalist", "品牌或引用異常、季度檢討", "整合 evidence 與決策歷史", "假設；頻率可能偏低"),
    ("餐飲品牌／門市", "品牌與地方型 discovery 查詢", "可作垂直案例與資料集", "應視為 Project；是否為直接 buyer 未證"),
], widths=[3.0, 4.4, 5.2, 3.6], font_size=8.6)
h2("3.3 為何不先主打所有商家")
p("一般商家通常購買結果，不購買量測方法；若 GeoCheck 無法證明可見度與來客或營收的關係，就很難要求商家長期訂閱。顧問與行銷團隊本來就需要報告、比較與決策紀錄，較能理解資料限制，也較可能形成多 Project 的固定工作流程。這仍是待驗證的 beachhead，不是已確認 ICP。")


h1("4. 現在做到哪裡", new_page=True)
p("以下只列 repo 與正式環境有記錄的狀態。『可運作』不等於『已證明有人願意付費』。更新基準為 2026 年 9 月 6 日。[S2]")
table(["項目", "狀態", "可提出的證據", "尚缺什麼"], [
    ("單站四題檢測", "已實作", "四個非品牌 discovery queries；自訂題可優先保留", "跨產業效度與長期穩定性"),
    ("AI 信任值 v1", "正式站已驗證", "65/35、有效分母、unknown、低樣本封頂規則已同步到 API／HTML／Markdown", "外部效度與使用者理解測試"),
    ("AI 供應商", "正式站已驗證", "Perplexity Sonar 與 DeepSeek V4 Flash 健康檢查成功", "持續可用性、成本與條款風險"),
    ("報告保存", "已啟用", "專用 Cloudflare D1；依網站、題組與 pipeline 版本辨識快取", "帳號、Project 權限與跨期資料模型"),
    ("GA4 漏斗", "程式已完成", "事件涵蓋 URL submit、分析、result_viewed、second_analysis、lead_submitted", "正式站 DebugView／Realtime 驗收與真實樣本"),
    ("白皮書研究治理", "規格與工具已建立", "人工題庫凍結、entity master、JSONL、雜湊與付費前 preflight", "人工核准樣本、正式付費 batch、外部審查"),
    ("Agent workspace", "概念與可行性評估", "已有 evidence graph／Decision Ledger／GeoBook 方向", "尚未完成外部使用者 decision loop"),
    ("商業化", "未驗證", "類別存在付費市場與競品收入自述", "GeoCheck 自有客戶、收入、WTP、留存、CAC、毛利"),
], widths=[2.7, 2.6, 6.2, 4.7], font_size=8.2)
label_box("最重要的現況判讀", "GeoCheck 已跨過『技術概念能否做出』，但還沒跨過『誰會持續使用並付費』。因此現在適合募的是驗證資源與育成支持，不適合用未經驗證的 ARR 故事募大額成長資金。", "FFF3E8")


h1("5. 核心創新與產品邊界", new_page=True)
h2("5.1 創新不等於多一個聊天視窗")
p("競品已普遍加入監測、建議、Agent 與工作流。GeoCheck 的差異化不能只是『用 AI 解釋 Dashboard』。可被防守的候選能力，是讓 Agent 受到證據圖譜約束：每個事實句要指向 observation，分母與 unknown 不得省略，建議要標信心與限制，使用者保留否決權。[S5–S6]")
h2("5.2 四層產品結構")
table(["層", "作用", "最小版本"], [
    ("Dashboard／Evidence layer", "保存答案、引用、有效性、版本與趨勢", "沿用現有報告與 D1，改成 Project 結構"),
    ("Geo Agent／Decision layer", "解釋變化、提出有限選項、承認資料不足", "只讀；固定 decision card schema"),
    ("GeoBooks／Workflow layer", "把人工驗證過的流程重跑", "先做一個『月度變化解釋』流程"),
    ("Decision Ledger", "保存接受／拒絕／延後理由，供下次驗證", "決策、證據、owner、驗證日與結果"),
], widths=[4.0, 6.2, 6.0])
h2("5.3 產品邊界")
bullet("前六個月不自動發文、不直接改網站、不做無限制 autonomous agent。")
bullet("不宣稱單一分數等於真實市場曝光、流量或營收。")
bullet("不為追求功能數量同時支援所有產業、所有引擎與 enterprise 權限。")
bullet("研究資料與客戶資料分開，只有在明確授權下才能用於改善流程。")


h1("6. 市場機會與產業判斷", new_page=True)
h2("6.1 市場存在，但 GeoCheck 的份額未知")
p("AI visibility／GEO monitoring 已有實際付費供給。Peec AI 曾公開宣布超過 US$4M ARR 與 1,300 個品牌／agency，Ahrefs、Semrush、Similarweb、Profound、Scrunch、Otterly、Hogiah 等也推出相鄰能力。[S5, S9] 這支持類別有預算，不代表台灣顧問會為 GeoCheck 付費，也不能用競品收入推估 GeoCheck 的 TAM。")
h2("6.2 市場階段")
p("本案判斷為快速成長與功能商品化同時發生。需求與供應正在增加，但 mention rate、citation、prompt monitoring 與泛用建議快速變成標準功能；大型 SEO suite 可以透過 bundle 降低小工具的獨立價值。GeoCheck 必須在窄工作流程、可追溯決策與在地資料品質上證明更高價值。")
h2("6.3 預算可能從哪裡來")
table(["可能預算來源", "購買理由", "風險"], [
    ("SEO／內容工具預算", "替代部分監測與報表工具", "大型 suite 已有相似模組"),
    ("顧問月報工時", "節省資料拼接、解釋與版本追蹤時間", "若每月只省很少時間，WTP 不高"),
    ("Agency retainer 毛利", "同一帳號服務多個長期客戶", "需要權限、模板、穩定交付"),
    ("研究／品牌情報預算", "需要可審計的 AI 答案與引用觀察", "市場較小，採購周期可能長"),
], widths=[4.0, 6.3, 5.9])


h1("7. 競爭與替代方案", new_page=True)
table(["競爭類型", "代表", "既有優勢", "GeoCheck 應對"], [
    ("Enterprise AI visibility", "Profound、Scrunch", "資料規模、團隊、enterprise sales", "不進 enterprise 功能戰"),
    ("專門 SaaS", "Peec、Otterly、Hogiah", "監測、報告、agency workflow、通路", "聚焦 evidence-first decision loop"),
    ("SEO Suite", "Ahrefs、Semrush、Similarweb", "既有資料、客群與 bundle", "不重建全套 SEO；以互補輸入整合"),
    ("泛用 Agent", "ChatGPT／Claude＋試算表", "便宜、彈性、使用習慣已形成", "用版本、證據與 decision ledger 降低 context 遺失"),
    ("不處理", "偶爾手查或不監測", "零成本、零導入負擔", "必須證明錯過決策的實際代價"),
], widths=[3.1, 3.6, 5.0, 4.5], font_size=8.5)
h2("7.1 Hogiah 的威脅")
p("Hogiah 在台灣語境、內容教育與在地入口上具明顯先行優勢，也已覆蓋許多 GeoCheck 原始 SaaS 構想。GeoCheck 不應跟著比引擎數、文章產量或功能清單；應先證明『有限資料下仍可給出可查核、可回頭驗證的判斷』能讓實作者節省高價值工時。")
h2("7.2 最強反方")
label_box("反方論點", "Ahrefs、Semrush 或現有 AI visibility 工具只要補上一個足夠好的 Agent 與 workflow，GeoCheck 的介面差異很快消失；如果目標使用者每月沒有高價值決策，Decision Ledger 也不會形成留存。", "FDECEC")
p("這個反方目前無法被產品功能反駁，只能由外部使用者的第二次使用、付款與決策紀錄反駁。")


h1("8. 建議目標客群與切入市場", new_page=True)
h2("8.1 建議 beachhead")
label_box("目標使用者假設", "台灣獨立 SEO／內容顧問與小型行銷工作室；已有 3–15 個長期服務客戶，固定製作月報或季度檢討，願意處理資料限制。", LIGHT)
p("這個選擇不是因為市場最大，而是因為一至二人的團隊較可能服務：同一使用者可帶入多個 Project，已有 recurring client relationship，不需要 GeoCheck 直接說服每家餐廳長期維護網站。")
h2("8.2 餐飲的角色")
p("餐飲適合作為第一個垂直資料包與公開案例：品牌多、地方 discovery query 明確、創辦人已有研究素材。它不應直接等同於初期 buyer。更合理的路徑是由顧問或工作室管理餐飲品牌 Project，GeoCheck 提供餐飲語境、entity 規則、題庫模板與證據流程。")
h2("8.3 何時擴張")
table(["擴張條件", "需要先看到的證據"], [
    ("從顧問擴到 in-house", "至少一種相同 decision loop 在不同組織仍重複發生"),
    ("從餐飲擴到其他垂直", "新垂直不需要大量例外 schema，且可重用 70% 以上流程"),
    ("從單一引擎擴到多引擎", "使用者願意為新增引擎資訊支付高於新增成本的價格"),
    ("從只讀 Agent 到執行", "建議正確性、人工批准與 rollback 已有足夠記錄"),
], widths=[5.0, 11.2])


h1("9. 商業模式", new_page=True)
h2("9.1 收費單位建議")
p("建議以活躍 Project 為主要 billing unit，而不是 Agent 對話次數。Project 對應一個品牌／網站、一組版本化觀測、有限監測與決策歷史；額外 provider 呼叫或更高頻率才按用量加購。這能讓價格與使用者管理的客戶數對齊。")
h2("9.2 從免費到付費的路徑")
table(["階段", "交付", "商業目的", "狀態"], [
    ("免費 Snapshot", "單站四題觀測＋有限證據", "讓使用者理解問題並留下下一步行為", "已有產品基礎"),
    ("Decision Loop 試用", "一次 Agent 解讀＋建立下一輪驗證", "測試是否願意投入資料與時間", "待開發／待驗證"),
    ("Practitioner Workspace", "3 個活躍 Project、月度 review、Decision Ledger", "形成訂閱與多客戶交付", "商業假設"),
    ("Studio Workspace", "更多 Project、模板、團隊 review", "擴大帳戶價值", "通過前兩階段後才做"),
], widths=[3.4, 5.6, 4.8, 2.4], font_size=8.5)
h2("9.3 定價不是現有成果")
p("正式定價尚未決定。可先用可退款預售測試單一 Project、固定監測與兩次 decision review。任何價格必須在實驗開始前凍結，並同時記錄 provider、儲存、人工 review、support 與獲客成本；在這些資料存在前，不計算虛假的 LTV／CAC 或毛利。")
label_box("投資人應追問", "付費是為了更多監測資料、節省報告工時、提高判斷可信度，還是取得餐飲垂直方法？只有付款訪談與實際使用能分辨。", "FFF3E8")


h1("10. 沒有人脈與 BD 能力時的獲客路徑", new_page=True)
h2("10.1 自媒體不是曝光策略，而是可量測漏斗")
p("創辦人可以用公開研究與真實案例取得第一批使用訊號，但發布篇數、觀看數與按讚數不算需求證據。每篇內容都要導向一個可觀察動作：提交網站、查看 evidence、建立下一輪提醒、授權最小 GSC CSV 或完成預售。")
h2("10.2 建議內容支柱")
table(["內容", "讀者得到什麼", "行為 CTA"], [
    ("台灣餐飲 AI 搜尋觀察", "知道 AI 答案引用誰、哪些證據不可靠", "提交品牌網址看自己的 evidence"),
    ("工具與方法拆解", "理解 mention rate、citation、unknown 的差異", "比較自己的既有工具輸出"),
    ("公開 decision log 案例", "看到建議如何連到下一輪驗證", "建立一個待驗證決策"),
    ("失敗與反例", "知道 GeoCheck 何時不該給答案", "提供反例或錯誤資料"),
], widths=[4.0, 7.1, 5.1])
h2("10.3 不靠熟人的驗證設計")
bullet("以搜尋、社群與案例頁取得至少 150 個符合目標內容的 sessions。")
bullet("追蹤非熟人提交、evidence 展開、下一輪提醒與第二次 result_viewed。")
bullet("只有合格使用者出現後才測預售；流量不足時只能判定 distribution 未驗證。")


h1("11. 六個月開發與驗證 Roadmap", new_page=True)
p("Roadmap 的順序由風險決定：先證明 Agent 不亂說，再證明外部使用者願意完成 decision loop，最後才證明付款與第二次使用。")
table(["月份", "核心工作", "交付物", "決策門檻"], [
    ("M0", "凍結範圍與測試資料", "20–30 個固定案例；evidence ID；錯誤／unknown fixtures", "可重現目前輸出"),
    ("M1", "Evidence Agent 正確性", "只讀 decision card；來源、信心、限制欄位", "事實句可追溯；無 unsupported claim"),
    ("M2", "單一 Project 與 Decision Ledger", "接受／拒絕／延後；owner；驗證日期", "創辦人可完成端到端回查"),
    ("M3", "公開案例與 self-serve CTA", "餐飲案例頁、提交入口、GA4 漏斗", "150 sessions 後至少 8 個外部提交、3 個深度行為"),
    ("M4", "第二輪驗證", "同一決策的 follow-up run 與差異解釋", "至少 3 人完成第二輪或提供資料承諾"),
    ("M5", "預售測試", "一 Project、固定監測、兩次 review 的可退款方案", "至少 2 筆預售，或 1 筆預售＋2 位第二輪資料授權"),
    ("M6", "投資／停止決策", "單位成本、使用記錄、質性回饋、下一版 scope", "達門檻才建最小訂閱 Workspace"),
], widths=[1.4, 4.2, 6.3, 4.3], font_size=8.1)
h2("11.1 通過後才開發")
bullet("多 Project 管理、報告模板與最小團隊 review。")
bullet("有限 GSC／GA4 匯入，用於對照而非宣稱因果校準。")
bullet("第二個 GeoBook 或第二個垂直 context package。")
page_break_before(h2("11.2 六個月內不要做"))
bullet("全球 prompt index、所有 AI 引擎、每日無限制監測。")
bullet("企業級 SSO、複雜席次權限、公開 API marketplace。")
bullet("自動寫文、發文、改網站與無人批准的執行 Agent。")
bullet("用模型微調取代資料結構、fixtures、規則與 evaluation。")
bullet("跨所有產業的 ranking portal 或單一『真實 ROI』魔法分數。")


h1("12. 成功指標、停損與可證偽假設")
table(["假設", "成功訊號", "失敗訊號", "失敗後處置"], [
    ("Agent 比 Dashboard 更有用", "使用者展開證據並建立待驗證決策", "只讀摘要、不採取任何下一步", "回到單次 evidence report"),
    ("有 recurring JTBD", "同一人於第二輪查看新的 result_viewed", "qualified users 無第二輪行為", "停止訂閱 monitor 假設"),
    ("有人願意付費", "達成 M5 預售門檻", "8 位 qualified preview 後無付款或資料承諾", "保留研究／作品用途"),
    ("自媒體能取客", "非熟人提交與目標角色相符", "只有同行／學生觀看，沒有 Project", "改受眾與入口，不擴產品"),
    ("成本可控", "每 Project 直接成本可量測且與價格相容", "反覆查詢、人工 review 或 support 失控", "降低頻率、縮工具或停止"),
], widths=[3.5, 4.6, 4.6, 3.5], font_size=8.1)
p("這些門檻是本提案的驗證建議，尚未成為正式產品 KPI。正式採用前應由創辦人確認並寫入 DECISION_LOG。")


h1("13. 團隊與執行能力", new_page=True)
h2("13.1 已展現的執行能力")
bullet("能把 AI 搜尋觀測拆成 provider、query、answer、citation、entity 與 denominator。")
bullet("已建立 unknown、低樣本封頂、資料雜湊、版本化快取與研究 preflight。")
bullet("已完成正式站 provider 健康檢查、D1 持久保存與前後端報告同步。")
bullet("能以競品反證調整方向，沒有把『競品沒有完整做到』當成藍海證據。")
h2("13.2 現階段團隊缺口")
table(["缺口", "影響", "六個月內的補法"], [
    ("沒有共同創辦團隊", "U-start 至少三人資格與執行風險", "找一位具 SEO／內容實務者與一位研究／設計或營運成員"),
    ("缺乏 BD／客戶網絡", "驗證速度與付費訪談不足", "用 self-serve 案例漏斗取得非熟人訊號；育成單位協助媒合"),
    ("缺乏產業 domain reviewer", "餐飲規則可能只反映創辦人想像", "邀請顧問或餐飲品牌行銷人審核 GeoBook"),
    ("一人承擔產品與研究", "scope 擴張、品質與節奏風險", "維持一個 Agent、一個 GeoBook、一個 ICP 假設"),
], widths=[3.4, 5.2, 7.6])
h2("13.3 U-start 資格提醒")
p("依 115 年度官方說明，U-start 團隊須至少三人，其中三分之二以上為近五學年度畢業生或大專校院在校生，並與設有育成單位的學校共同申請。[S7–S8] 以目前已知的一人創辦狀態，尚不具完整申請條件；這是組隊與育成合作缺口，不是計畫內容可以代替的文件問題。")


h1("14. 財務規劃與資金用途", new_page=True)
h2("14.1 現況")
p("GeoCheck 尚無可驗證營收、客戶數、CAC、LTV、續約率或完整單位成本。因此本章提供的是里程碑預算框架，不是財務預測。正式募資前要以 provider 帳單、雲端費用、人工 review 時數、內容取得成本與預售結果更新。")
h2("14.2 六個月驗證預算框架（建議）")
table(["用途", "占比建議", "支出目的", "撥款門檻"], [
    ("產品與雲端", "30%", "provider、D1／運算、監測與必要開發工具", "只支付固定案例與外部 Project 所需"),
    ("市場驗證", "25%", "合格受測者回饋、預售與訪談行政", "有明確對象與紀錄格式"),
    ("內容與案例", "20%", "公開研究、案例頁、內容剪輯與基本素材", "每項內容綁定 CTA"),
    ("研究與專業審查", "15%", "方法、domain 與資安／隱私檢查", "審查項目與修正紀錄可交付"),
    ("法務、會計與預備", "10%", "公司、合約、隱私與不可預期費用", "需要時才動用"),
], widths=[3.5, 2.6, 6.5, 3.6], font_size=8.3)
h2("14.3 資金策略建議")
p("目前較適合申請育成補助、競賽資源或小額 milestone-based angel tranche，用來完成 M1–M5 的驗證。尚不建議提出以規模擴張為目的的大額 angel round，因為需求、付款、重複使用與 acquisition 都未成立。實際募資金額與股權條件應在完成成本盤點後由創辦人決定，本提案不代填估值或投資條款。")
h2("14.4 U-start 經費語境")
p("115 年度官方說明載明，第一階段通過者可取得合計新臺幣 50 萬元，其中 35 萬元為團隊創業基本開辦費、15 萬元為學校育成輔導費；第二階段另有績優獎助。[S8] 本預算表可作為經費合理性的初稿，但仍需依下一年度正式須知、學校會計規則與核銷科目重編。")


h1("15. 投資判斷與資金請求", new_page=True)
h2("15.1 建議投資結論")
label_box("結論：有條件進入驗證", "GeoCheck 有可信的技術與研究基礎，也位於已有付費供給的市場；但目前缺乏 GeoCheck 自有需求、付款與留存證據。適合以六個月里程碑支持驗證，不適合把產品完成度誤當成商業成立。", "FFF3E8")
h2("15.2 三個支持理由")
number("已有可運作且重視分母、unknown、entity 與版本的 measurement layer，降低從零建產品的風險。")
number("AI visibility 類別已有付費市場與多家供應者，題目不是純概念；GeoCheck 無須先證明類別存在。")
number("半 Agent＋Dashboard 可用有限、只讀、可追溯的方式開發，符合一至二人團隊的成本邊界。")
h2("15.3 三個反對理由")
number("尚無目標客群訪談、付費、第二次使用或 acquisition 證據。")
number("Agent、monitoring、GSC workflow 已被競品與大型 SEO suite 快速商品化。")
number("創辦人缺乏團隊與 BD 網絡；自媒體通路尚未證實能觸達有資料權限的 buyer。")
h2("15.4 投資人可採用的條件")
table(["條件", "里程碑", "未達成時"], [
    ("產品可信", "M1 evidence agent 正確性 gate 通過", "停止擴充 Agent capability"),
    ("需求存在", "M3 外部提交與深度行為達門檻", "調整受眾與問題，不補功能"),
    ("留存可能", "M4 出現第二輪 decision loop", "停止 subscription thesis"),
    ("付款存在", "M5 預售／資料承諾達門檻", "不進 full SaaS build"),
    ("成本相容", "可算每 Project 直接成本與人工負擔", "降低監測或停止"),
], widths=[3.0, 8.0, 5.2])
page_break_before(h2("15.5 本輪需要的非金錢資源"))
bullet("學校育成單位與至少兩位互補成員，以滿足 U-start 申請與執行需求。")
bullet("5–10 位真正負責 SEO／內容月報的外部受測者。")
bullet("餐飲或地方服務業 domain reviewer，以及能挑戰方法的研究讀者。")


h1("16. U-start 評選項目對照")
p("115 年第一階段官方評分以產業與市場／商業模式 40%、學校育成能力 20%、團隊執行與未來發展 15%、目標與效益 15%、財務規劃 10% 組成。[S7]")
table(["官方評選項目", "比重", "本提案對應", "目前缺口"], [
    ("產業、市場與商業模式", "40%", "第 6–10 章：市場、競爭、ICP、billing unit、獲客", "缺第一手 buyer／WTP／單位經濟"),
    ("學校育成輔導能力", "20%", "第 13、15 章列出需要的媒合、domain 與商業輔導", "尚未確定合作學校與育成計畫"),
    ("團隊執行力與未來發展", "15%", "第 4、5、11、13 章列出既有系統與 milestone", "目前一人；角色配置未完成"),
    ("目標與預期效益", "15%", "第 1、3、12 章說明問題、價值與可量測成功／失敗", "商業效益仍須外部驗證"),
    ("財務規劃", "10%", "第 14 章提供驗證導向經費框架", "缺實際報價、核銷科目與現金流"),
], widths=[4.0, 1.6, 6.5, 4.1], font_size=8.2)
h2("16.1 目前參賽成熟度")
table(["面向", "判斷", "原因"], [
    ("題目與技術", "可提案", "已有產品、方法與正式環境證據"),
    ("市場與商業", "需補證", "仍是二手市場證據與未驗證 ICP"),
    ("團隊與資格", "目前不足", "至少三人與學校育成合作尚未完成"),
    ("財務", "可做初稿", "有用途框架，沒有正式金額與報價"),
    ("整體", "適合準備下一輪", "先組隊並完成 M1–M3，提案可信度會明顯提高"),
], widths=[4.0, 3.0, 9.2])


h1("17. 主要風險與治理", new_page=True)
table(["風險", "可能性／衝擊", "早期訊號", "處理方式"], [
    ("大型 suite bundle Agent", "高／高", "既有客戶可在原工具完成相同工作", "聚焦 evidence schema、Decision Ledger 與窄 GeoBook"),
    ("使用頻率不足", "中高／高", "使用者只看一次 Snapshot", "只在第二輪成立後建 subscription"),
    ("Agent 過度解讀", "中／高", "無來源建議、誤讀 unknown、否決率高", "只讀、固定 schema、人工批准、eval fixtures"),
    ("供應商與成本", "中／高", "API 失敗、條款改變、每 Project 成本上升", "快取、硬上限、provider 記錄、停止規則"),
    ("資料與隱私", "中／高", "GSC／GA4 權限疑慮或資料混用", "最小權限、明確授權、研究／客戶資料隔離"),
    ("單人 scope 爆炸", "高／中高", "多 Agent、多垂直、多引擎同時進行", "一個 Agent、一個 workflow、一個 beachhead"),
], widths=[3.5, 2.7, 5.4, 4.6], font_size=8.0)
h2("17.1 治理原則")
p("所有重大定位、目標客群、商業模式、定價、評分與產品優先序，仍由創辦人拍板並寫入 DECISION_LOG。這份提案提供建議，不將未確認方向寫成正式決策。[S4]")


h1("18. 證據來源與名詞", new_page=True)
h2("18.1 證據標記")
table(["標記", "意思"], [
    ("已驗證事實", "repo、正式環境、官方規則或可追溯來源直接支持"),
    ("觀察結果", "在特定時間、網站、query 或測試條件下看到的結果"),
    ("推論", "由多項證據推得，但仍可能有替代解釋"),
    ("建議", "為降低風險提出的行動順序，尚未成為正式決策"),
    ("假設", "可被未來測試推翻的客群、價值、價格或通路主張"),
    ("未知", "目前沒有足夠資料，不能填 0 或用合理故事補上"),
], widths=[3.8, 12.4])
h2("18.2 來源清單")
sources = [
    ("S1", "GeoCheck Project Charter", "目的、研究主軸、邊界與未驗證客群。", "repo: PROJECT_CHARTER.md"),
    ("S2", "GeoCheck Current State", "正式站、D1、GA4 與未完成門檻。", "repo: CURRENT_STATE.md"),
    ("S3", "AI Trust Index v1.0.0", "65/35、分母、unknown 與版本邊界。", "repo: AI_TRUST_INDEX_V1.md"),
    ("S4", "GeoCheck Decision Log", "正式決策與商業模式未定狀態。", "repo: DECISION_LOG.md"),
    ("S5", "競品證偽與產品路線建議", "市場、競爭、Modify 建議與限制。", "repo: competitive-falsification/REPORT.md"),
    ("S6", "Agentic Workspace Idea Check", "Evidence Workspace 與驗證門檻。", "repo: agentic-workspace-idea-check.md"),
    ("S7", "U-start 115 年評選及審查方式", "第一階段五項評選與比重。", "https://ustart.yda.gov.tw/p/16-1000-526.php?Lang=zh-tw"),
    ("S8", "教育部青年發展署 U-start 計畫頁", "參加資格、期程、補助與育成合作。", "https://www.yda.gov.tw/plan.aspx?p=2013"),
    ("S9", "Peec AI Series A announcement", "供應商自報 ARR 與品牌／agency 數，僅作類別市場證據。", "https://peec.ai/blog/we-raised-21m-series-a-to-help-brands-win-in-ai-search"),
]
table(["ID", "來源", "本提案用途", "位置"], sources, widths=[1.2, 3.8, 6.2, 5.0], font_size=7.0)
h2("18.3 本提案仍不知道的事")
bullet("誰是 GeoCheck 第一個穩定付款 buyer，以及他從哪一筆預算支付。")
bullet("一個 Project 的真實直接成本、人工 review 時數與可接受價格。")
bullet("是否每月都有足以改變決策的 AI visibility 變化。")
bullet("自媒體能否帶來有網站與資料權限的合格使用者。")
bullet("Decision Ledger 與 GeoBook 是否能形成足夠的切換成本。")


h1("附錄 A：創辦人對外簡報的 90 秒版本")
p("GeoCheck 是一個面向生成式搜尋的證據型工作空間。品牌現在開始在意 AI 是否提到自己、引用哪些網站，但多數工具停在監測圖表；顧問仍要自己判斷變化是否可信、下一步做什麼，以及下個月如何驗證。")
p("我們已完成單站四題觀測、AI 信任值、unknown 與有效分母規則、正式站供應商驗證、版本化報告保存和研究治理。下一步不是再堆一個 Dashboard，而是加入受限 Geo Agent：它只能根據已保存的 evidence 解釋變化，每個判斷都有來源、信心與限制，使用者可以否決，下一輪再回查決策。")
p("首批客群假設是有長期客戶的 SEO／內容顧問與小型工作室，餐飲作為第一個垂直案例。未來六個月先驗證三件事：Agent 是否能不亂說、外部使用者是否完成第二輪 decision loop、以及是否有人願意預售。達標才建立訂閱 Workspace；未達標就停止擴大 SaaS，保留研究與一次性檢測價值。")


h1("附錄 B：下一步行動清單", new_page=True)
table(["優先序", "行動", "完成定義", "負責角色"], [
    ("1", "確認是否以顧問／小型工作室作為 beachhead", "創辦人拍板並記入 DECISION_LOG", "創辦人"),
    ("2", "盤點 U-start 組隊資格與育成單位", "三人角色草案、學校窗口與下一年度時程", "創辦人＋學校"),
    ("3", "完成 Evidence Agent correctness gate", "固定案例通過引用、unknown 與 unsupported-claim 檢查", "產品／研究"),
    ("4", "建立一個餐飲公開案例漏斗", "案例、提交、evidence 展開、提醒與 GA4 全流程", "產品／內容"),
    ("5", "執行外部使用與預售實驗", "依 M3–M5 門檻得到可判斷結果", "創辦人／營運"),
], widths=[1.6, 5.0, 6.5, 3.1], font_size=8.4)
p("本文件截至 2026 年 9 月 7 日。官方競賽資格、補助金額與申請期程應在實際申請時重新核對當年度最新須知。", italic=True)


# Document metadata
doc.core_properties.title = "GeoCheck 創業提案書｜投資人暨 U-start 評審閱讀版"
doc.core_properties.subject = "GeoCheck 專案說明、現況、商業路徑、Roadmap 與投資判斷"
doc.core_properties.author = "GeoCheck"
doc.core_properties.keywords = "GeoCheck, GEO, AI Visibility, U-start, Angel Investment, Business Plan"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
